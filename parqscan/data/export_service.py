from __future__ import annotations

import csv
import json
import math
import re
from pathlib import Path
from typing import Callable

import pyarrow as pa
import pyarrow.parquet as pq
from openpyxl import Workbook

from parqscan.constants import EXPORT_BATCH_ROWS
from parqscan.data.parquet_document import ParquetDocument
from parqscan.utils.serialization import json_safe

Progress = Callable[[int, int], None]
Cancelled = Callable[[], bool]


def _python_rows(batch: pa.RecordBatch) -> list[dict]:
    return [{name: json_safe(value) for name, value in row.items()} for row in batch.to_pylist()]


def _flat_value(value):
    if isinstance(value, (dict, list, tuple)):
        return json.dumps(json_safe(value), ensure_ascii=False)
    return value


def export_csv(document: ParquetDocument, target: str, progress: Progress, cancelled: Cancelled) -> None:
    path = Path(target)
    processed = 0
    with path.open("w", newline="", encoding="utf-8-sig") as stream:
        writer = csv.DictWriter(stream, fieldnames=document.schema.names)
        writer.writeheader()
        for batch in document.iter_batches_from(0, EXPORT_BATCH_ROWS):
            if cancelled():
                break
            for row in _python_rows(batch):
                writer.writerow({key: _flat_value(value) for key, value in row.items()})
            processed += batch.num_rows
            progress(processed, document.total_rows)


def export_json(document: ParquetDocument, target: str, progress: Progress, cancelled: Cancelled) -> None:
    processed = 0
    first = True
    with Path(target).open("w", encoding="utf-8") as stream:
        stream.write("[\n")
        for batch in document.iter_batches_from(0, EXPORT_BATCH_ROWS):
            if cancelled():
                break
            for row in _python_rows(batch):
                if not first:
                    stream.write(",\n")
                json.dump(row, stream, ensure_ascii=False)
                first = False
            processed += batch.num_rows
            progress(processed, document.total_rows)
        stream.write("\n]\n")


def export_excel(document: ParquetDocument, target: str, progress: Progress, cancelled: Cancelled) -> None:
    workbook = Workbook(write_only=True)
    try:
        sheet_index = 1
        sheet = workbook.create_sheet(f"Data {sheet_index}")
        sheet.append(document.schema.names)
        row_in_sheet = 1
        processed = 0
        for batch in document.iter_batches_from(0, EXPORT_BATCH_ROWS):
            if cancelled():
                break
            for row in _python_rows(batch):
                if row_in_sheet >= 1_000_000:
                    sheet_index += 1
                    sheet = workbook.create_sheet(f"Data {sheet_index}")
                    sheet.append(document.schema.names)
                    row_in_sheet = 1
                sheet.append([_flat_value(row.get(name)) for name in document.schema.names])
                row_in_sheet += 1
            processed += batch.num_rows
            progress(processed, document.total_rows)
        workbook.save(target)
    finally:
        workbook.close()


def export_markdown(document: ParquetDocument, target: str, progress: Progress, cancelled: Cancelled) -> None:
    processed = 0
    with Path(target).open("w", encoding="utf-8") as stream:
        stream.write("| " + " | ".join(document.schema.names) + " |\n")
        stream.write("| " + " | ".join("---" for _ in document.schema.names) + " |\n")
        for batch in document.iter_batches_from(0, EXPORT_BATCH_ROWS):
            if cancelled():
                break
            for row in _python_rows(batch):
                values = [str(_flat_value(row.get(name))).replace("|", "\\|").replace("\n", "<br>") for name in document.schema.names]
                stream.write("| " + " | ".join(values) + " |\n")
            processed += batch.num_rows
            progress(processed, document.total_rows)


def _sql_type(data_type: pa.DataType) -> str:
    if pa.types.is_integer(data_type):
        return "BIGINT"
    if pa.types.is_floating(data_type) or pa.types.is_decimal(data_type):
        return "DOUBLE"
    if pa.types.is_boolean(data_type):
        return "BOOLEAN"
    if pa.types.is_date(data_type) or pa.types.is_timestamp(data_type):
        return "TIMESTAMP"
    if pa.types.is_binary(data_type) or pa.types.is_large_binary(data_type) or pa.types.is_fixed_size_binary(data_type):
        return "BLOB"
    return "TEXT"


def _sql_literal(value) -> str:
    if value is None:
        return "NULL"
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if isinstance(value, float) and (math.isnan(value) or math.isinf(value)):
            return "NULL"
        return str(value)
    text = str(_flat_value(value)).replace("'", "''")
    return f"'{text}'"


def export_sql(document: ParquetDocument, target: str, progress: Progress, cancelled: Cancelled) -> None:
    table_name = re.sub(r"[^a-zA-Z0-9_]", "_", Path(document.path).stem) or "parquet_data"
    processed = 0
    with Path(target).open("w", encoding="utf-8") as stream:
        definitions = [f'  "{field.name}" {_sql_type(field.type)}' for field in document.schema]
        stream.write(f'CREATE TABLE "{table_name}" (\n' + ",\n".join(definitions) + "\n);\n\n")
        column_sql = ", ".join(f'"{name}"' for name in document.schema.names)
        for batch in document.iter_batches_from(0, EXPORT_BATCH_ROWS):
            if cancelled():
                break
            for row in _python_rows(batch):
                values = ", ".join(_sql_literal(row.get(name)) for name in document.schema.names)
                stream.write(f'INSERT INTO "{table_name}" ({column_sql}) VALUES ({values});\n')
            processed += batch.num_rows
            progress(processed, document.total_rows)


def export_parquet_fragment(
    document: ParquetDocument,
    target: str,
    start_row: int,
    end_row: int,
    columns: list[str],
    progress: Progress,
    cancelled: Cancelled,
) -> None:
    row_count = max(0, end_row - start_row + 1)
    writer: pq.ParquetWriter | None = None
    processed = 0
    try:
        for batch in document.iter_batches_from(start_row, EXPORT_BATCH_ROWS, columns=columns):
            if cancelled() or processed >= row_count:
                break
            remaining = row_count - processed
            if batch.num_rows > remaining:
                batch = batch.slice(0, remaining)
            table = pa.Table.from_batches([batch])
            if writer is None:
                writer = pq.ParquetWriter(target, table.schema)
            writer.write_table(table)
            processed += batch.num_rows
            progress(processed, row_count)
    finally:
        if writer is not None:
            writer.close()
        elif not cancelled():
            pq.write_table(pa.Table.from_batches([], schema=document.selected_schema(columns)), target)


EXPORT_HANDLERS = {
    "csv": export_csv,
    "json": export_json,
    "excel": export_excel,
    "markdown": export_markdown,
    "sql": export_sql,
}
